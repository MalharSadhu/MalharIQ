import io
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


def create_executive_deck(
    biz_report: dict,
    tech_report: dict,
    ml_results: dict,
    split_counts: dict,
    task: str,
) -> io.BytesIO:
  """Generates a structured, McKinsey-style consulting PowerPoint presentation."""
  prs = Presentation()
  prs.slide_width = Inches(13.33)  # 16:9 widescreen format
  prs.slide_height = Inches(7.5)

  blank_slide_layout = prs.slide_layouts[6]

  # Brand Color Palette (Navy / Slate / Blue)
  NAVY = RGBColor(15, 23, 42)
  BLUE = RGBColor(37, 99, 235)
  SLATE = RGBColor(100, 116, 139)
  WHITE = RGBColor(255, 255, 255)
  CARD_BG = RGBColor(248, 250, 252)

  def add_header(slide, title_text: str, category_text: str = "EXECUTIVE BRIEF"):
    header_box = slide.shapes.add_textbox(
        Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.0)
    )
    tf = header_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    p_cat = tf.paragraphs[0]
    p_cat.text = category_text.upper()
    p_cat.font.size = Pt(10)
    p_cat.font.bold = True
    p_cat.font.color.rgb = BLUE

    p_title = tf.add_paragraph()
    p_title.text = title_text
    p_title.font.size = Pt(22)
    p_title.font.bold = True
    p_title.font.color.rgb = NAVY

  # --------------------------------------------------------------------------
  # SLIDE 1: Title & Mandate
  # --------------------------------------------------------------------------
  s1 = prs.slides.add_slide(blank_slide_layout)
  tbox = s1.shapes.add_textbox(
      Inches(1.0), Inches(2.2), Inches(11.3), Inches(3.0)
  )
  tf1 = tbox.text_frame
  tf1.word_wrap = True

  p1 = tf1.paragraphs[0]
  p1.text = "MALHAR IQ | ANALYTICS ADVISORY"
  p1.font.size = Pt(12)
  p1.font.bold = True
  p1.font.color.rgb = BLUE

  p2 = tf1.add_paragraph()
  p2.text = f"Strategic Machine Learning Mandate:\n{task}"
  p2.font.size = Pt(32)
  p2.font.bold = True
  p2.font.color.rgb = NAVY

  p3 = tf1.add_paragraph()
  p3.text = (
      "Autonomous Quantitative Benchmark, Holdout Validation & Governance"
      " Audit"
  )
  p3.font.size = Pt(14)
  p3.font.color.rgb = SLATE

  # --------------------------------------------------------------------------
  # SLIDE 2: Executive Verdict & Recommendations
  # --------------------------------------------------------------------------
  s2 = prs.slides.add_slide(blank_slide_layout)
  add_header(s2, "Executive Findings & Strategic Action Items")

  # Left Box: Verdict
  box_left = s2.shapes.add_textbox(
      Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8)
  )
  tf_l = box_left.text_frame
  tf_l.word_wrap = True

  p_vl = tf_l.paragraphs[0]
  p_vl.text = "Executive Verdict"
  p_vl.font.size = Pt(16)
  p_vl.font.bold = True
  p_vl.font.color.rgb = NAVY

  p_vt = tf_l.add_paragraph()
  p_vt.text = biz_report.get("executive_verdict", "No verdict generated.")
  p_vt.font.size = Pt(12)
  p_vt.font.color.rgb = SLATE

  tf_l.add_paragraph()  # spacer
  p_sl = tf_l.add_paragraph()
  p_sl.text = "Key Operational Insights:"
  p_sl.font.size = Pt(13)
  p_sl.font.bold = True
  p_sl.font.color.rgb = NAVY

  for ins in biz_report.get("segment_insights", [])[:3]:
    p = tf_l.add_paragraph()
    p.text = f"• {ins}"
    p.font.size = Pt(11)
    p.font.color.rgb = SLATE

  # Right Box: Action Items
  box_right = s2.shapes.add_textbox(
      Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.8)
  )
  tf_r = box_right.text_frame
  tf_r.word_wrap = True

  p_rl = tf_r.paragraphs[0]
  p_rl.text = "Strategic Recommendations"
  p_rl.font.size = Pt(16)
  p_rl.font.bold = True
  p_rl.font.color.rgb = BLUE

  for rec in biz_report.get("strategic_recommendations", [])[:4]:
    p = tf_r.add_paragraph()
    p.text = f"• {rec}"
    p.font.size = Pt(12)
    p.font.color.rgb = NAVY

  # --------------------------------------------------------------------------
  # SLIDE 3: Architecture Benchmark & Data Splits
  # --------------------------------------------------------------------------
  s3 = prs.slides.add_slide(blank_slide_layout)
  add_header(s3, "Architecture Benchmark & 70/15/15 Validation Protocol")

  box_m = s3.shapes.add_textbox(
      Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0)
  )
  tf_m = box_m.text_frame
  tf_m.word_wrap = True

  p_champ = tf_m.paragraphs[0]
  p_champ.text = (
      f"Selected Champion Architecture: {ml_results.get('best_algorithm', 'N/A')}"
  )
  p_champ.font.size = Pt(16)
  p_champ.font.bold = True
  p_champ.font.color.rgb = BLUE

  tot_rows = sum(split_counts.values()) if split_counts else 1
  p_split = tf_m.add_paragraph()
  p_split.text = (
      f"Dataset Partitioning: Train Set = {split_counts.get('train', 0):,} rows"
      f" ({split_counts.get('train', 0)/tot_rows*100:.1f}%) | "
      f"Validation Set = {split_counts.get('val', 0):,} rows"
      f" ({split_counts.get('val', 0)/tot_rows*100:.1f}%) | "
      f"Holdout Test Set = {split_counts.get('test', 0):,} rows"
      f" ({split_counts.get('test', 0)/tot_rows*100:.1f}%)"
  )
  p_split.font.size = Pt(12)
  p_split.font.color.rgb = SLATE

  tf_m.add_paragraph()  # spacer
  p_arch = tf_m.add_paragraph()
  p_arch.text = "Engineering Architecture Summary:"
  p_arch.font.size = Pt(13)
  p_arch.font.bold = True
  p_arch.font.color.rgb = NAVY

  p_arch_txt = tf_m.add_paragraph()
  p_arch_txt.text = tech_report.get(
      "architecture_summary", "Architecture evaluation complete."
  )
  p_arch_txt.font.size = Pt(12)
  p_arch_txt.font.color.rgb = SLATE

  # --------------------------------------------------------------------------
  # SLIDE 4: Risk, Bias & Governance Audit
  # --------------------------------------------------------------------------
  s4 = prs.slides.add_slide(blank_slide_layout)
  add_header(s4, "Governance, Bias Parity & Risk Audit")

  box_gov = s4.shapes.add_textbox(
      Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0)
  )
  tf_g = box_gov.text_frame
  tf_g.word_wrap = True

  p_gov_h = tf_g.paragraphs[0]
  p_gov_h.text = "Demographic Parity & Four-Fifths Compliance (EEOC Standard)"
  p_gov_h.font.size = Pt(16)
  p_gov_h.font.bold = True
  p_gov_h.font.color.rgb = NAVY

  p_gov_txt = tf_g.add_paragraph()
  p_gov_txt.text = (
      "Pipeline applies deterministic disparate impact auditing. Any cohort"
      " receiving under 80% relative favorable selection rate triggers an"
      " active operational warning."
  )
  p_gov_txt.font.size = Pt(12)
  p_gov_txt.font.color.rgb = SLATE

  tf_g.add_paragraph()  # spacer
  p_risk_h = tf_g.add_paragraph()
  p_risk_h.text = "Key Operational & Model Risk Factors:"
  p_risk_h.font.size = Pt(14)
  p_risk_h.font.bold = True
  p_risk_h.font.color.rgb = BLUE

  for rk in biz_report.get("potential_business_risks", []):
    p = tf_g.add_paragraph()
    p.text = f"• {rk}"
    p.font.size = Pt(12)
    p.font.color.rgb = SLATE

  # Save presentation to memory buffer
  pptx_buffer = io.BytesIO()
  prs.save(pptx_buffer)
  pptx_buffer.seek(0)
  return pptx_buffer